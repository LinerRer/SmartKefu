"""文本解析器：把 PDF / DOCX / TXT / MD 统一转成纯文本。

设计原则：
1. 解析结果为后续"递归分块"服务——必须保留段落边界（空行），
   不能把全文压成一行，否则分块器没有语义边界可用。
2. 不认识的扩展名直接拒绝，宁可报错也不猜。
"""
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def extract_text(file_path: str) -> str:
    ext = Path(file_path).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"不支持的文件类型: {ext}")
    if ext == ".pdf":
        text = _extract_pdf(file_path)
    elif ext == ".docx":
        text = _extract_docx(file_path)
    else:
        text = _extract_plain_text(file_path)
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _extract_pdf(path: str) -> str:
    """PyMuPDF 逐页提取文本，页与页之间补空行。"""
    import pymupdf

    parts = []
    with pymupdf.open(path) as doc:
        for page in doc:
            parts.append(page.get_text())
    return "\n\n".join(parts)


def _extract_docx(path: str) -> str:
    """python-docx 按文档顺序提取段落和表格，段间/行间保留空行。

    只提取 paragraphs 会漏掉表格里的内容（题库/产品参数常见），
    所以按 body 子元素顺序遍历，段落和表格交替输出。
    """
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(path)
    # 建立 paragraph / table 对象的 xml 元素 -> 对象 映射，用于按顺序取对象
    para_map = {p._element: p for p in doc.paragraphs}
    table_map = {t._element: t for t in doc.tables}

    parts: list[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn('w:p'):
            p = para_map.get(child)
            if p and p.text.strip():
                parts.append(p.text.strip())
        elif child.tag == qn('w:tbl'):
            t = table_map.get(child)
            if t:
                for row in t.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(' | '.join(cells))
    return "\n\n".join(parts)


def _extract_plain_text(path: str) -> str:
    """先按字节读，再依次尝试编码，全失败才报错。"""
    raw = Path(path).read_bytes()
    for enc in ("utf-8", "gbk", "utf-8-sig"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"无法识别文件编码: {path}")
